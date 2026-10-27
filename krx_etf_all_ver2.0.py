#!/usr/bin/env python3
"""
KRX ETF '기타 ETF' 추출 스크립트 - [클라우드(GitHub Actions)용 ver2.0]

흐름: KRX 로그인(ID/PW, 헤드리스) -> [13104] ETF 전종목 기본정보 조회 -> 기타 ETF 필터링/엑셀 생성 -> 메일 발송

  국내주식형 ETF = 기초시장분류 '국내' AND 기초자산분류 '주식' AND 추적배수 '일반'
  기타 ETF       = 위 조건에 해당하지 않는 모든 ETF
  출력 컬럼      = Code(단축코드 + ' KP') / ISIN(표준코드) / Name(영문종목명)

ver1 -> ver2 변경점
  - 사람이 브라우저에서 로그인하던 방식 -> KRX 아이디/비밀번호로 자동 로그인 (매 실행마다 새로 로그인)
  - 인증정보는 코드에 넣지 않고 환경변수(GitHub Secrets)로만 받는다
      KRX_ID, KRX_PW, SMTP_USER, SMTP_PASS   (선택: MAIL_TO, SMTP_PROVIDER, SMTP_HOST, SMTP_PORT, OUT_DIR)
  - 결과 폴더는 OUT_DIR (기본 output), 날짜는 한국시간(KST) 기준
  - 로그인/조회 실패 시 OUT_DIR/debug/ 에 스크린샷·HTML을 남김 (워크플로에서 아티팩트로 업로드)

설치
  pip install playwright pandas openpyxl
  python -m playwright install --with-deps chromium

실행
  KRX_ID=... KRX_PW=... SMTP_USER=... SMTP_PASS=... python krx_etf_all_ver2.0.py
  python krx_etf_all_ver2.0.py --no-mail            # 엑셀만 생성
  python krx_etf_all_ver2.0.py --headed --no-mail   # (PC 테스트용) 브라우저 화면을 띄워서 로그인 과정 확인
  python krx_etf_all_ver2.0.py --file X.xlsx --no-mail   # 이미 받은 KRX 엑셀로 필터/엑셀만 테스트
"""
import argparse
import datetime as dt
import json
import os
import re
import smtplib
import ssl
import sys
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

# ============================== 설정 ==============================
DEFAULT_TO = "youngwook0002.kim@samsung.com"
PROVIDERS = {"gmail": ("smtp.gmail.com", 465), "naver": ("smtp.naver.com", 465)}

OUT_DIR = Path(os.getenv("OUT_DIR") or "output").resolve()
OUT_DIR.mkdir(parents=True, exist_ok=True)
DEBUG_DIR = OUT_DIR / "debug"

LOGIN_URL = os.getenv("KRX_LOGIN_URL") or "https://data.krx.co.kr/contents/MDC/COMS/client/MDCCOMS001.cmd"
HOME_URL = "https://data.krx.co.kr/contents/MDC/MDI/mdiLoader/index.cmd?menuId=MDC0201"
KRX_URL = "https://data.krx.co.kr/comm/bldAttendant/getJsonData.cmd"
KRX_BLD = "dbms/MDC/STAT/standard/MDCSTAT04601"  # ETF 전종목 기본정보
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")

# 로그인 화면 입력칸 후보 (화면 구조가 바뀌면 여기만 고치면 됨)
ID_SELECTORS = ["#mbrId", "input[name='mbrId']", "input[id*='mbrId']",
                "input[name*='id' i]:not([type='hidden']):not([type='password'])",
                "form input[type='text']"]
PW_SELECTORS = ["#pw", "input[name='pw']", "input[type='password']"]
BTN_SELECTORS = ["button:has-text('로그인')", "a:has-text('로그인')",
                 "input[type='submit']", "input[type='button'][value*='로그인']"]

FIELD_MAP = {  # KRX 응답 필드 -> 한글 컬럼
    "ISU_SRT_CD": "단축코드",
    "ISU_CD": "표준코드",
    "ISU_ENG_NM": "영문종목명",
    "IDX_MKT_CLSS_NM": "기초시장분류",
    "IDX_ASST_CLSS_NM": "기초자산분류",
}
OUT_COLS = {"단축코드": "Code", "표준코드": "ISIN", "영문종목명": "Name"}


def today_kst() -> dt.date:
    try:
        from zoneinfo import ZoneInfo
        return dt.datetime.now(ZoneInfo("Asia/Seoul")).date()
    except Exception:  # noqa: BLE001
        return dt.date.today()


FILE_NAME = f"[Samsung] Korea other ETF list_{today_kst():%y%m%d}.xlsx"


def env(name: str) -> str:
    v = (os.getenv(name) or "").strip()
    if not v:
        sys.exit(f"환경변수 {name} 가 없습니다. (GitHub: Settings > Secrets and variables > Actions)")
    return v


# ============================ 1) 로그인 + 조회 ============================
def save_debug(page, tag: str) -> None:
    """실패 원인 확인용 스크린샷/HTML 저장 (비밀번호는 화면에 마스킹되어 남지 않음)."""
    try:
        DEBUG_DIR.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(DEBUG_DIR / f"{tag}.png"), full_page=True)
        (DEBUG_DIR / f"{tag}.html").write_text(page.content(), encoding="utf-8")
        print(f"  [debug] {DEBUG_DIR / tag}.png / .html 저장")
    except Exception as e:  # noqa: BLE001
        print(f"  [debug] 저장 실패: {e}")


def first_visible(page, selectors, timeout=15000):
    """후보 셀렉터 중 화면에 보이는 첫 번째 요소 (iframe 안쪽까지 탐색)."""
    import time
    deadline = time.time() + timeout / 1000
    while time.time() < deadline:
        for frame in page.frames:
            for sel in selectors:
                loc = frame.locator(sel)
                try:
                    for i in range(min(loc.count(), 5)):
                        if loc.nth(i).is_visible():
                            return loc.nth(i)
                except Exception:  # noqa: BLE001
                    continue
        page.wait_for_timeout(500)
    return None


def login(ctx, user: str, pw: str) -> None:
    page = ctx.new_page()
    # '다른 곳에서 로그인 중입니다. 계속하시겠습니까?' 같은 확인창은 자동으로 '확인'
    page.on("dialog", lambda d: (print(f"  [로그인] 알림창: {d.message.strip()[:80]}"), d.accept()))
    print("  [로그인] 로그인 화면 이동")
    page.goto(LOGIN_URL, timeout=90000, wait_until="domcontentloaded")

    id_box = first_visible(page, ID_SELECTORS)
    pw_box = first_visible(page, PW_SELECTORS)
    if not (id_box and pw_box):
        save_debug(page, "login_form_not_found")
        sys.exit("로그인 입력칸을 찾지 못했습니다. output/debug/login_form_not_found.png 를 확인하세요.")
    id_box.fill(user)
    pw_box.fill(pw)

    btn = first_visible(page, BTN_SELECTORS, timeout=3000)
    if btn:
        btn.click()
    else:
        pw_box.press("Enter")
    try:
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:  # noqa: BLE001
        pass
    page.wait_for_timeout(2000)

    # 로그인 실패 메시지가 화면에 남아 있으면 바로 중단 (잘못된 비밀번호로 반복 시도해 계정 잠기는 것 방지)
    body = page.inner_text("body")[:5000] if page.locator("body").count() else ""
    if re.search(r"비밀번호.{0,10}(일치하지|틀렸|오류)|존재하지 않는 (아이디|회원)|계정이 잠", body):
        save_debug(page, "login_rejected")
        sys.exit("KRX 로그인이 거부되었습니다(아이디/비밀번호 확인). 재시도하지 않고 종료합니다.")
    page.close()


def fetch_basic(ctx) -> pd.DataFrame | None:
    page = ctx.new_page()
    try:
        page.goto(HOME_URL, timeout=90000, wait_until="domcontentloaded")
        resp = ctx.request.post(
            KRX_URL,
            form={"bld": KRX_BLD, "locale": "ko_KR", "csvxls_isNo": "false"},
            headers={"Referer": HOME_URL, "X-Requested-With": "XMLHttpRequest"},
            timeout=90000,
        )
        status, text = resp.status, resp.text()
    except Exception as e:  # noqa: BLE001
        print(f"  조회 중 오류: {str(e).splitlines()[0]}")
        save_debug(page, "fetch_error")
        return None
    if status != 200:
        print(f"  KRX 응답 status={status} (로그인 안 된 상태일 가능성)")
        save_debug(page, f"fetch_status_{status}")
        return None
    try:
        rows = json.loads(text).get("output", [])
    except ValueError:
        print("  JSON이 아닌 응답 (로그아웃 상태 가능성)")
        DEBUG_DIR.mkdir(parents=True, exist_ok=True)
        (DEBUG_DIR / "fetch_not_json.txt").write_text(text[:20000], encoding="utf-8")
        return None
    finally:
        page.close()
    return pd.DataFrame(rows) if rows else None


def get_etf_basic(headed: bool = False) -> pd.DataFrame:
    from playwright.sync_api import sync_playwright

    user, pw = env("KRX_ID"), env("KRX_PW")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not headed)
        ctx = browser.new_context(user_agent=UA, locale="ko-KR", timezone_id="Asia/Seoul")
        try:
            login(ctx, user, pw)
            df = fetch_basic(ctx)
            if df is None:   # 로그인 직후 세션 반영이 늦는 경우 대비 1회만 재시도
                print("  1회 재로그인 후 재시도")
                login(ctx, user, pw)
                df = fetch_basic(ctx)
        finally:
            browser.close()
    if df is None:
        sys.exit("로그인 후에도 데이터를 받지 못했습니다. output/debug/ 파일을 확인하세요.")
    print(f"  ETF 전종목 기본정보 {len(df)}건 수신")
    return df


# ============================ 2) 필터링 / 엑셀 ============================
def normalize_krx(raw: pd.DataFrame) -> pd.DataFrame:
    """KRX 응답 -> 한글 컬럼. 추적배수는 필드명이 불확실해 값으로 탐지."""
    df = raw.copy()
    mult_col = next(
        (c for c in df.columns
         if (df[c].astype(str) == "일반").any()
         and df[c].astype(str).str.contains("레버리지|인버스").any()),
        None,
    )
    if mult_col is None:
        sys.exit(f"추적배수(일반/레버리지/인버스) 컬럼을 찾지 못했습니다. 응답 컬럼: {list(df.columns)}")
    rename = {k: v for k, v in FIELD_MAP.items() if k in df.columns}
    rename[mult_col] = "추적배수"
    return df.rename(columns=rename)


def extract_other_etf(df: pd.DataFrame) -> pd.DataFrame:
    need = ["단축코드", "표준코드", "영문종목명", "기초시장분류", "기초자산분류", "추적배수"]
    missing = [c for c in need if c not in df.columns]
    if missing:
        sys.exit(f"필요한 컬럼이 없습니다: {missing}\n현재 컬럼: {list(df.columns)}")
    df = df.copy()
    for c in need:
        df[c] = df[c].astype(str).str.strip()
    is_domestic_equity = (
        (df["기초시장분류"] == "국내") & (df["기초자산분류"] == "주식") & (df["추적배수"] == "일반")
    )
    out = df.loc[~is_domestic_equity, list(OUT_COLS)].rename(columns=OUT_COLS)
    out["Code"] = out["Code"] + " KP"  # 6자리 단축코드 뒤에 ' KP' 추가 (예: 0193M0 KP)
    return out.drop_duplicates("ISIN").reset_index(drop=True)


def save_excel(df: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="OtherETF")
        ws = w.sheets["OtherETF"]
        for col, width in zip("ABC", (10, 16, 70)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        for c in ws[1]:  # 헤더 행 서식
            c.font = Font(name="Arial", size=10, bold=True, color="000000")
            c.fill = PatternFill("solid", fgColor="D9D9D9")
            c.alignment = Alignment(horizontal="center", vertical="center")


# ============================ 3) 메일 발송 ============================
def build_message(xlsx: Path, n_other: int, n_total: int) -> EmailMessage:
    user = env("SMTP_USER")
    to = [x.strip() for x in re.split(r"[,;\s]+", os.getenv("MAIL_TO") or DEFAULT_TO) if x.strip()]
    msg = EmailMessage()
    msg["From"] = os.getenv("MAIL_FROM") or user
    msg["To"] = ", ".join(to)
    msg["Subject"] = f"[Samsung] Korea other ETF list_{today_kst():%y%m%d}"
    msg.set_content(
        f"KRX ETF 전종목 기본정보 기준 기타 ETF {n_other}개 (전체 {n_total}개) 목록을 첨부합니다.\n"
        "\n"
        "[필터링 룰]\n"
        "- 국내주식형 ETF = 기초시장분류 '국내' + 기초자산분류 '주식' + 추적배수 '일반' (3개 조건 모두 충족)\n"
        "- 위 국내주식형 ETF를 제외한 나머지는 모두 기타 ETF로 분류\n"
        "  (해외, 국내&해외, 채권, 혼합자산, 원자재, 부동산, 통화, 레버리지/인버스 등)\n"
    )
    msg.add_attachment(
        xlsx.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=xlsx.name,
    )
    return msg


def send_mail(msg: EmailMessage) -> None:
    provider = (os.getenv("SMTP_PROVIDER") or "gmail").lower()
    host, port = PROVIDERS.get(provider, (None, None))
    host = os.getenv("SMTP_HOST") or host
    port = int(os.getenv("SMTP_PORT") or port or 465)
    if not host:
        sys.exit("SMTP_PROVIDER(gmail|naver) 또는 SMTP_HOST 환경변수를 설정하세요.")
    user = env("SMTP_USER")
    pw = env("SMTP_PASS").replace(" ", "")
    ctxt = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctxt, timeout=30) as s:
            s.login(user, pw)
            s.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=30) as s:
            s.starttls(context=ctxt)
            s.login(user, pw)
            s.send_message(msg)


# ============================== main ==============================
def main():
    try:
        sys.stdout.reconfigure(line_buffering=True)
    except Exception:  # noqa: BLE001
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-mail", action="store_true", help="메일 발송 없이 엑셀만 생성")
    ap.add_argument("--headed", action="store_true", help="(PC 테스트용) 브라우저 화면을 띄워서 실행")
    ap.add_argument("--file", help="(테스트용) 이미 받은 KRX 엑셀로 실행 -- 로그인/조회 생략")
    ap.add_argument("--out", default=str(OUT_DIR / FILE_NAME))
    args = ap.parse_args()

    if args.file:
        df = pd.read_excel(args.file, dtype=str)
    else:
        df = normalize_krx(get_etf_basic(headed=args.headed))

    others = extract_other_etf(df)
    out = Path(args.out)
    save_excel(others, out)
    print(f"전체 {len(df)}개 중 기타 ETF {len(others)}개 -> {out}")

    if args.no_mail:
        return
    msg = build_message(out, len(others), len(df))
    send_mail(msg)
    print(f"메일 발송 완료 -> {msg['To']}")


if __name__ == "__main__":
    main()
